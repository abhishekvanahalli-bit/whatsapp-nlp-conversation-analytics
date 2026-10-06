"""Build the 8-slide MSc presentation from the frozen report, the fictional sample screenshots and Figure 2.

Run from the project root:  python presentation/build_presentation.py
All screenshots are of the fictional sample chat (data/sample_chat.txt). Numbers are copied from the final report.
"""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "presentation" / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------- palette (the dashboard's)
BG, CARD, BORDER = "EDF1F7", "FFFFFF", "CDD6E4"
INK, MUTED = "172033", "525D70"
BLUE, TEAL, PURPLE, ORANGE, GREEN = "1E5AE0", "08766A", "6240C8", "BF5410", "1D7A3B"
TINT = {BLUE: "E1EBFD", TEAL: "D9F1ED", PURPLE: "EBE5FA", ORANGE: "FCE6D4", GREEN: "DCF1E3"}
FONT = "Calibri"


def rgb(h): return RGBColor.from_string(h)


# ---------------------------------------------------------------- image crops (fictional sample only)
def crop(src, box, name):
    im = Image.open(ROOT / src).convert("RGB").crop(box)
    out = ASSETS / name; im.save(out, optimize=True); return out, im.size


overview_top, ov_size = crop("docs/images/overview.png", (0, 0, 1440, 725), "overview_top.png")
conv, cv_size = crop("docs/images/conversation.png", (250, 0, 1440, 1075), "conversation_main.png")
tfidf, tf_size = crop("docs/images/nlp_tfidf.png", (250, 0, 1440, 1235), "tfidf_main.png")
patt, pt_size = crop("docs/images/patterns.png", (250, 0, 1440, 900), "patterns_main.png")
pipe = ROOT / "report" / "figure2_pipeline.png"
pipe_size = Image.open(pipe).size

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


# ---------------------------------------------------------------- helpers
def new_slide(notes):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    s.notes_slide.notes_text_frame.text = notes
    return s


def shadow(shape):
    sp = shape._element.spPr
    for e in sp.findall(qn("a:effectLst")): sp.remove(e)
    eff = etree.SubElement(sp, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="76200", dist="25400", dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="101828"); etree.SubElement(clr, qn("a:alpha"), val="12000")


def card(s, x, y, w, h, fill=CARD, radius=0.04, line=BORDER, shade=True):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.adjustments[0] = radius
    r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    if line: r.line.color.rgb = rgb(line); r.line.width = Pt(0.75)
    else: r.line.fill.background()
    if shade: shadow(r)
    return r


def text(s, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, italic=False):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {bold,color,size,italic}) tuples."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.space_after = Pt(spacing)
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run(); r.text = t; f = r.font
            f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def title(s, t, sub=None):
    text(s, 0.6, 0.42, 10.5, 0.7, t, size=36, bold=True)
    if sub: text(s, 0.6, 1.08, 11.5, 0.4, sub, size=16, color=MUTED)


def footer(s, n):
    text(s, 0.6, 7.05, 9, 0.25, "WhatsApp NLP & Conversation Analytics Dashboard  |  MSc Data Science  |  Abhishek Vanahalli", size=11, color=MUTED)
    text(s, 12.0, 7.05, 0.75, 0.25, str(n), size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def pic(s, path, size, x, y, w=None, h=None, frame=True):
    """Place an image inside a white card, preserving aspect ratio; returns (x,y,w,h) of the image."""
    iw, ih = size
    if h is not None and w is None: w = h * iw / ih
    if w is not None and h is None: h = w * ih / iw
    if frame: card(s, x - 0.1, y - 0.1, w + 0.2, h + 0.2, radius=0.025)
    s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return x, y, w, h


def chip(s, x, y, w, label, color=BLUE, h=0.34, size=12):
    r = card(s, x, y, w, h, fill=TINT[color], radius=0.5, line=None, shade=False)
    tf = r.text_frame; tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    rr = p.add_run(); rr.text = label; rr.font.name = FONT; rr.font.size = Pt(size); rr.font.bold = True; rr.font.color.rgb = rgb(color)


def circle_num(s, x, y, n, color=BLUE, d=0.42):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = rgb(color); c.line.fill.background()
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); r.font.name = FONT; r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = rgb("FFFFFF")


def kpi(s, x, y, w, h, label, value, note, color=BLUE):
    card(s, x, y, w, h)
    text(s, x + 0.18, y + 0.14, w - 0.36, 0.26, label.upper(), size=11, bold=True, color=MUTED)
    text(s, x + 0.18, y + 0.38, w - 0.36, 0.5, value, size=28, bold=True, color=INK)
    text(s, x + 0.18, y + h - 0.32, w - 0.36, 0.26, note, size=12, color=MUTED)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + w - 0.46), Inches(y + 0.14), Inches(0.26), Inches(0.26))
    dot.fill.solid(); dot.fill.fore_color.rgb = rgb(TINT[color]); dot.line.fill.background()


# ================================================================ 1 Title
s = new_slide("Introduce the project in one sentence: a reproducible, privacy-aware pipeline and dashboard that describes the language and activity structure of an exported WhatsApp chat. Mention that the public demonstration uses a fictional chat.")
chip(s, 0.7, 0.75, 1.9, "MSc DATA SCIENCE", BLUE)
text(s, 0.7, 1.35, 5.9, 2.3, ["WhatsApp NLP &", "Conversation Analytics", "Dashboard"], size=40, bold=True)
text(s, 0.7, 3.65, 5.7, 0.9, "A privacy-aware pipeline and interactive dashboard for describing the language and activity structure of exported WhatsApp chats", size=16, color=MUTED)
card(s, 0.7, 4.85, 5.7, 1.75)
text(s, 0.95, 5.02, 5.3, 1.5, [
    [("Abhishek Vanahalli", {"bold": True, "size": 18})],
    [("CHRIST (Deemed to be University)", {"size": 15})],
    [("Supervisor: Himani Sivaraman, Faculty, CHRIST (Deemed to be University)", {"size": 14, "color": MUTED})],
    [("10/10/2026", {"size": 14, "color": MUTED})]], spacing=3)
pic(s, overview_top, ov_size, 6.95, 1.95, w=5.9)
text(s, 6.95, 5.2, 5.9, 0.3, "Overview page, fictional sample chat", size=12, color=MUTED, italic=True)

# ================================================================ 2 Problem & objectives
s = new_slide("Problem: an exported chat is a raw text file mixing messages, system events, media placeholders and several scripts. Objectives: parse robustly, apply NLP (n-grams, TF-IDF), describe script usage and key terms, extract conversation-level patterns, and present everything in a privacy-conscious dashboard.")
title(s, "Problem & Objectives", "Why a raw chat export needs a careful, explainable analysis pipeline")
card(s, 0.6, 1.75, 5.5, 4.95)
text(s, 0.9, 1.95, 4.9, 0.4, "PROBLEM", size=13, bold=True, color=BLUE)
text(s, 0.9, 2.35, 4.9, 0.9, "An exported WhatsApp chat is a raw text file. It is hard to see directly:", size=17, bold=True)
probs = ["Message patterns", "Frequently used terms", "Activity over time", "Repeated content", "Writing-system usage", "Conversation composition"]
for i, t_ in enumerate(probs):
    yy = 3.35 + i * 0.52
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), Inches(yy + 0.07), Inches(0.16), Inches(0.16))
    d.fill.solid(); d.fill.fore_color.rgb = rgb(BLUE if i % 2 == 0 else TEAL); d.line.fill.background()
    text(s, 1.3, yy, 4.5, 0.4, t_, size=17)
card(s, 6.4, 1.75, 6.33, 4.95)
text(s, 6.7, 1.95, 5.6, 0.4, "OBJECTIVES", size=13, bold=True, color=BLUE)
objs = ["Build a robust WhatsApp parser and preprocessing pipeline",
        "Apply NLP techniques: n-grams and TF-IDF",
        "Analyse Script Mix and key conversational terms",
        "Extract activity and repetition patterns",
        "Present results in a privacy-conscious interactive dashboard"]
for i, o in enumerate(objs):
    yy = 2.5 + i * 0.84
    circle_num(s, 6.7, yy, i + 1, [BLUE, TEAL, PURPLE, ORANGE, GREEN][i])
    text(s, 7.3, yy - 0.04, 5.2, 0.7, o, size=17, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)

# ================================================================ 3 Data & methodology
s = new_slide("Walk down the pipeline: validate, parse, preprocess and classify rows, then the four text analyses and conversation analytics with 28 reconciliation checks, then privacy filtering, the results bundle and the dashboard. Stress that the public demonstration uses a fictional chat so nothing private is exposed.")
title(s, "Data & Methodology", "One supported export format, one reproducible pipeline")
ph = 5.2
card(s, 0.6, 1.65, 4.6, 5.25)
s.shapes.add_picture(str(pipe), Inches(0.6 + (4.6 - ph * pipe_size[0] / pipe_size[1]) / 2), Inches(1.75), height=Inches(ph - 0.2))
text(s, 5.55, 1.75, 7.2, 0.3, "WHAT HAPPENS TO AN UPLOADED CHAT", size=13, bold=True, color=BLUE)
steps = [("Validate and parse", "Bracketed .txt export; header, system and user rows; multi-line messages kept together", BLUE),
         ("Preprocess and classify", "Text, media or deleted; links and phone numbers masked; Unicode-aware tokens", TEAL),
         ("Analyse", "N-grams, TF-IDF, Script Mix, Key Terms and conversation analytics (28 reconciliation checks)", PURPLE),
         ("Filter and present", "Privacy filter, anonymised results bundle, Streamlit dashboard", ORANGE)]
for i, (h_, d_, c_) in enumerate(steps):
    yy = 2.2 + i * 0.98
    card(s, 5.55, yy, 7.18, 0.86, radius=0.12)
    circle_num(s, 5.75, yy + 0.22, i + 1, c_, d=0.42)
    text(s, 6.4, yy + 0.09, 6.1, 0.3, h_, size=16, bold=True)
    text(s, 6.4, yy + 0.4, 6.15, 0.4, d_, size=13, color=MUTED)
card(s, 5.55, 6.18, 7.18, 0.72, fill=TINT[BLUE], radius=0.15, line=None, shade=False)
text(s, 5.8, 6.18, 6.7, 0.72, "Public demonstration uses a fictional WhatsApp-style dataset for privacy and reproducibility.", size=15, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 3)

# ================================================================ 4 NLP analysis
s = new_slide("Four analyses. N-grams are counted per message over all messages and over unique texts, so repetition is visible. TF-IDF weights distinctiveness within the chat and was validated against scikit-learn. Script Mix describes writing systems and is NOT language identification. Key Terms rank terms by the number of distinct messages containing them.")
title(s, "NLP Analysis")
cards4 = [("N-grams", BLUE, ["Unigrams", "Bigrams", "Trigrams", "All messages vs unique texts"]),
          ("TF-IDF", PURPLE, ["Term distinctiveness", "Primary + unique-text variants", "Validated against scikit-learn"]),
          ("Script Mix", TEAL, ["Latin", "Devanagari", "Mixed-script", "No-alphabetic"]),
          ("Key Terms", ORANGE, ["Message coverage", "Total occurrences", "Repeated-message effect"])]
for i, (h_, c_, items) in enumerate(cards4):
    cx, cy = 0.6 + (i % 2) * 3.7, 1.35 + (i // 2) * 2.5
    card(s, cx, cy, 3.5, 2.3)
    chip(s, cx + 0.2, cy + 0.2, 1.7, h_.upper(), c_)
    for j, it in enumerate(items):
        text(s, cx + 0.25, cy + 0.72 + j * 0.37, 3.1, 0.34, it, size=15)
card(s, 0.6, 6.4, 7.2, 0.52, fill=TINT[TEAL], radius=0.2, line=None, shade=False)
text(s, 0.85, 6.4, 6.8, 0.52, "Script Mix is a writing-system analysis, not language identification.", size=15, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
iw4 = 4.3
x_, y_, w_, h_ = pic(s, tfidf, tf_size, 8.5, 1.5, w=iw4)
text(s, 8.5, 1.5 + h_ + 0.25, 4.5, 0.3, "TF-IDF tab, fictional sample chat", size=12, color=MUTED, italic=True)
footer(s, 4)

# ================================================================ 5 Conversation analytics
s = new_slide("All numbers are from the fictional sample chat. 168 rows: 1 header, 7 system, 160 user rows. Of the user rows, 143 text, 12 media, 5 deleted; 12 participants. Activity on 13 calendar days with a peak of 15 rows on 2026-03-14. 99 of 143 text messages repeat an earlier text, leaving 44 unique texts in 35 repeated templates.")
title(s, "Conversation Analytics", "Fictional sample dataset: 12 invented participants")
grid = [("Total rows", "168", "parsed", BLUE), ("User rows", "160", "of 168 rows", BLUE), ("Text messages", "143", "analysed", TEAL),
        ("Media rows", "12", "placeholders", ORANGE), ("Deleted rows", "5", "placeholders", ORANGE), ("Participants", "12", "pseudonymised", PURPLE),
        ("Active days", "13", "calendar days", GREEN), ("Peak activity", "15 rows", "on 2026-03-14", GREEN), ("Repeated texts", "99 / 143", "text messages", ORANGE),
        ("Unique texts", "44", "after masking", TEAL), ("Templates", "35", "repeated", PURPLE)]
kw, kh, gx, gy = 1.95, 1.12, 0.12, 0.12
for i, (l, v, n, c_) in enumerate(grid):
    kpi(s, 0.6 + (i % 3) * (kw + gx), 1.75 + (i // 3) * (kh + gy), kw, kh, l, v, n, c_)
card(s, 0.6 + 2 * (kw + gx), 1.75 + 3 * (kh + gy), kw, kh, fill=TINT[BLUE], radius=0.1, line=None, shade=False)
text(s, 0.6 + 2 * (kw + gx) + 0.15, 1.75 + 3 * (kh + gy), kw - 0.3, kh, "Fictional sample dataset", size=15, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)
ph_ = 4.6
x_, y_, w_, h_ = pic(s, conv, cv_size, 7.1, 1.85, h=ph_)
text(s, 7.1, 1.85 + ph_ + 0.25, 5.6, 0.3, "Conversation page, fictional sample chat", size=12, color=MUTED, italic=True)
footer(s, 5)

# ================================================================ 6 Dashboard
s = new_slide("The dashboard is a Streamlit and Plotly prototype. Upload a chat, it is processed in a temporary folder for the session, and four areas show the results with privacy filtering applied. The screenshot is the Overview page for the fictional sample chat.")
title(s, "Interactive Dashboard", "Streamlit + Plotly academic prototype; upload-driven, session-based analysis")
iw = 8.6
x_, y_, w_, h_ = pic(s, overview_top, ov_size, 0.7, 1.75, w=iw)
text(s, 0.7, 1.75 + h_ + 0.22, iw, 0.3, "Overview page loaded with the fictional sample chat (12 invented participants)", size=12, color=MUTED, italic=True)
for i, (h_t, d_t, c_) in enumerate([("Overview", "key figures, daily activity, composition, key findings", BLUE),
                                    ("Conversation", "activity, composition, repetition, hour and weekday tables", TEAL),
                                    ("NLP Insights", "six tabs: n-grams, TF-IDF, Script Mix, Key Terms", PURPLE),
                                    ("Patterns", "repetition summary and anonymous templates", ORANGE)]):
    yy = 1.65 + i * 1.12
    card(s, 9.75, yy, 3.0, 1.0, radius=0.1)
    text(s, 9.95, yy + 0.1, 2.7, 0.3, h_t, size=16, bold=True, color=c_)
    text(s, 9.95, yy + 0.42, 2.65, 0.55, d_t, size=12, color=MUTED)
chip(s, 0.7, 6.6, 2.1, "Privacy-conscious display", TEAL)
chip(s, 2.95, 6.6, 1.9, "Anonymised download", BLUE)
chip(s, 5.0, 6.6, 2.0, "Optional theme image", PURPLE)
footer(s, 6)

# ================================================================ 7 Validation & limitations
s = new_slide("Validation: 304 tests pass in a clean clone (308 locally, including 4 tests that need a private chat and are excluded from the public repository), 28 reconciliation checks, TF-IDF agrees with scikit-learn to within about 1.1e-16, and the committed outputs are reproduced. Privacy: fictional public dataset, private chat excluded, repository scanned. Limitations are stated openly: one supported format, Script Mix is not language identification, the privacy filter is not complete PII detection, and results are not population-level.")
title(s, "Validation & Limitations", "What was checked, and what the results do not show")
cols = [("VALIDATION", BLUE, ["304 tests passed in a clean clone", "308 tests locally (4 are local-only)", "28 reconciliation checks", "TF-IDF vs scikit-learn: max difference about 1.1 × 10⁻¹⁶", "Committed outputs reproduced"]),
        ("PRIVACY", TEAL, ["Fictional public dataset", "Private chat excluded from the repository", "Privacy and display filtering", "Repository privacy scan", "No private chat results shown publicly"]),
        ("LIMITATIONS", ORANGE, ["Supported export format is limited", "Android / no-bracket formats unsupported", "Script Mix is not language identification", "Privacy filter is not complete PII detection", "Public results are not population-level findings"])]
for i, (h_, c_, items) in enumerate(cols):
    cx = 0.6 + i * 4.15
    card(s, cx, 1.75, 3.95, 4.95)
    chip(s, cx + 0.25, 1.98, 1.7, h_, c_)
    for j, it in enumerate(items):
        yy = 2.58 + j * 0.8
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.28), Inches(yy + 0.09), Inches(0.15), Inches(0.15))
        d.fill.solid(); d.fill.fore_color.rgb = rgb(c_); d.line.fill.background()
        text(s, cx + 0.6, yy, 3.15, 0.72, it, size=15)
footer(s, 7)

# ================================================================ 8 Conclusion & future work
s = new_slide("Conclusion: a reproducible, tested pipeline and dashboard for exported chats, with limits stated. The fictional dataset allows safe demonstration; the numbers show pipeline behaviour and reproducibility, not general WhatsApp use. Future work comes from the report: wider export formats, larger real evaluation, better multilingual handling, stronger personal-data detection. Then invite questions.")
title(s, "Conclusion & Future Work")
card(s, 0.6, 1.4, 6.1, 4.2)
text(s, 0.9, 1.6, 5.5, 0.35, "CONCLUSION", size=13, bold=True, color=BLUE)
text(s, 0.9, 2.0, 5.5, 1.0, "A reproducible pipeline for analysing WhatsApp-style exported conversations, with its limits stated.", size=17, bold=True)
for i, it in enumerate(["Preprocessing, n-grams and TF-IDF", "Script Mix and Key Terms", "Conversation analytics and an interactive dashboard", "Validation shows reproducibility and robustness", "A fictional public dataset enables safe demonstration"]):
    yy = 3.05 + i * 0.52
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), Inches(yy + 0.09), Inches(0.15), Inches(0.15))
    d.fill.solid(); d.fill.fore_color.rgb = rgb(BLUE if i % 2 == 0 else TEAL); d.line.fill.background()
    text(s, 1.28, yy, 5.2, 0.45, it, size=15)
card(s, 6.95, 1.4, 5.78, 4.2)
text(s, 7.25, 1.6, 5.2, 0.35, "FUTURE WORK", size=13, bold=True, color=PURPLE)
for i, it in enumerate(["Broader WhatsApp export-format support", "Larger, real-world evaluation (group and iOS exports)", "Richer multilingual analysis and stopword handling", "Stronger personal-data detection", "More robust large-scale TF-IDF evaluation"]):
    yy = 2.2 + i * 0.65
    circle_num(s, 7.25, yy, i + 1, PURPLE, d=0.38)
    text(s, 7.8, yy - 0.04, 4.7, 0.5, it, size=15, anchor=MSO_ANCHOR.MIDDLE)
card(s, 0.6, 5.85, 12.13, 0.95, fill=BLUE, radius=0.12, line=None)
text(s, 0.9, 5.85, 6, 0.95, "Thank you", size=30, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
text(s, 7.0, 5.85, 5.5, 0.95, "Questions?", size=30, bold=True, color="FFFFFF", align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 8)

prs.core_properties.title = "WhatsApp NLP & Conversation Analytics Dashboard"
prs.core_properties.author = "Abhishek Vanahalli"
out = ROOT / "presentation" / "WhatsApp_NLP_Conversation_Analytics_Presentation.pptx"
prs.save(out); print("saved", out)
